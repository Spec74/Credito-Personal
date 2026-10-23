# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x11, 0x48, 0x85)
NEGRO = RGBColor(0x1F, 0x29, 0x37)
GRIS = RGBColor(0x59, 0x5F, 0x6B)
OUT = Path(__file__).with_name("Preguntas-para-el-gerente-Crediconfiable.docx")


def set_run(run, size=11, bold=False, color=NEGRO, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def p(doc, text, *, size=11, bold=False, color=NEGRO, italic=False, space_after=8, space_before=0, align="left"):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_run(para.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return para


def pregunta(doc, numero, texto):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    para.paragraph_format.space_after = Pt(4)
    r1 = para.add_run(f"{numero}.  ")
    set_run(r1, size=12, bold=True, color=AZUL)
    r2 = para.add_run(texto)
    set_run(r2, size=12, bold=False, color=NEGRO)

    raya = doc.add_paragraph()
    raya.paragraph_format.space_before = Pt(2)
    raya.paragraph_format.space_after = Pt(2)
    set_run(raya.add_run("Respuesta:"), size=10, italic=True, color=GRIS)
    for _ in range(3):
        linea = doc.add_paragraph()
        linea.paragraph_format.space_after = Pt(0)
        linea.paragraph_format.space_before = Pt(0)
        set_run(linea.add_run("________________________________________________________________"))


def main():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    p(doc, "INVERSIONES CREDICONFIABLE S.A.C.", size=11, bold=True, color=AZUL, space_after=2, align="center")
    p(doc, "Ayacucho, 2026", size=11, color=GRIS, space_after=14, align="center")
    p(
        doc,
        "Preguntas para el gerente",
        size=20,
        bold=True,
        color=AZUL,
        space_after=6,
        align="center",
    )
    p(
        doc,
        "Trabajo de tesis universitaria — se responde por escrito",
        size=11,
        italic=True,
        color=GRIS,
        space_after=16,
        align="center",
    )

    p(doc, "Estimado gerente:", size=12, space_after=8)
    p(
        doc,
        "Le escribo estas preguntas para un trabajo de la universidad. Quiero entender cómo usted "
        "mira el negocio: si los clientes pagan, cómo cierra el mes y qué le gustaría ver más fácil. "
        "No hay respuestas correctas. Puede escribir corto, como le salga.",
        size=12,
        space_after=8,
        align="justify",
    )
    p(
        doc,
        "Si una pregunta no le aplica, ponga “no”. Si quiere, puede contestarme por WhatsApp, "
        "pregunta por pregunta.",
        size=12,
        space_after=16,
        align="justify",
    )

    p(doc, "Cobranza", size=14, bold=True, color=AZUL, space_before=4, space_after=4)
    pregunta(
        doc,
        1,
        "Cuando usted quiere saber si los clientes están pagando, ¿qué hace primero? ¿Entra al sistema, abre un Excel, o le pregunta a alguien?",
    )
    pregunta(
        doc,
        2,
        "Ese dato, ¿lo ve en soles? Por ejemplo: “nos deben 80 mil”.",
    )
    pregunta(
        doc,
        3,
        "¿También lo ve en porcentaje? Por ejemplo: “está al 18”.",
    )
    pregunta(
        doc,
        4,
        "¿Eso lo mira de toda la empresa junto, o lo separa por oficina (Huamanga, Pichari, Huanta, Andahuaylas)?",
    )
    pregunta(
        doc,
        5,
        "¿Y lo mira por analista? ¿Quién está cobrando bien y quién no?",
    )
    pregunta(
        doc,
        6,
        "Si un cliente se atrasa un día, ¿ya le preocupa o espera unos días? ¿Cuántos días espera?",
    )
    pregunta(
        doc,
        7,
        "Acuérdese de la última vez que una oficina o un analista no estaba cobrando bien. ¿Cómo se enteró usted? ¿Qué hizo: lo llamó, fue a ver, o lo dejó?",
    )

    p(doc, "El fin de mes", size=14, bold=True, color=AZUL, space_before=16, space_after=4)
    pregunta(
        doc,
        8,
        "Cuando se acaba el mes, ¿qué números quiere tener sí o sí? Por ejemplo: cuánto hay prestado, cuánto no han pagado, cuánto se cobró.",
    )
    pregunta(
        doc,
        9,
        "Esos números del mes, ¿quién los saca? ¿Usted, el administrador o contabilidad?",
    )
    pregunta(
        doc,
        10,
        "¿En cuánto tiempo los tiene? ¿El mismo día, a la semana, o más?",
    )
    pregunta(
        doc,
        11,
        "¿Esos números salen del sistema o alguien los pasa a un Excel?",
    )
    pregunta(
        doc,
        12,
        "Si usan un Excel a mano: ¿para qué le sirve? ¿Qué no le da el sistema?",
    )

    p(doc, "Oficinas y metas", size=14, bold=True, color=AZUL, space_before=16, space_after=4)
    pregunta(
        doc,
        13,
        "¿Usted les dice a las oficinas o a los analistas “este mes tienen que cobrar tanto” o “prestar tanto”? Si sí, ¿lo anota en algún lado o lo tiene en la cabeza?",
    )
    pregunta(
        doc,
        14,
        "Cuando decide, ¿necesita ver las cuatro oficinas juntas o a veces solo una (por ejemplo solo Pichari)?",
    )
    pregunta(
        doc,
        15,
        "Eso lo mira todos los días, una vez a la semana, o recién cuando cierra el mes?",
    )
    pregunta(
        doc,
        16,
        "¿Lo mira en la computadora de la oficina o también en el celular?",
    )
    pregunta(
        doc,
        17,
        "Además de lo prestado y lo que no pagan, ¿le hace falta ver la caja del día o los gastos de la empresa? ¿O con eso le alcanza?",
    )

    p(doc, "Para el trabajo de la universidad", size=14, bold=True, color=AZUL, space_before=16, space_after=4)
    pregunta(
        doc,
        18,
        "¿Me permite estar un rato el día que arman los números del mes? Solo miraría y anotaría cuánto demoran. Ese día no se cambia nada.",
    )
    pregunta(
        doc,
        19,
        "Si esto le ayuda, ¿qué le gustaría poder hacer que hoy le cuesta? Una o dos cosas, como usted lo dice. Por ejemplo: ver las oficinas en un solo lugar, no esperar el Excel, o saber el mismo día quién no está pagando.",
    )
    pregunta(
        doc,
        20,
        "En el trabajo de la universidad saldrían totales (“la empresa prestó tanto”, “no pagaron tanto”), sin nombres ni DNI de clientes. ¿Está de acuerdo en que se usen así y en firmar un papel de eso? ¿Hay algo que no deba salir de la empresa (sueldos, lista de clientes, otra cosa)?",
    )

    p(doc, "", size=6, space_after=8)
    p(
        doc,
        "Gracias por su tiempo. Puede devolver este mismo archivo con las respuestas, o escribirme por WhatsApp.",
        size=12,
        align="justify",
        space_after=16,
    )
    p(doc, "Atentamente,", size=12, space_after=2)
    p(doc, "________________________________", size=12, space_after=0)
    p(doc, "(Nombre de quien envía)", size=10, italic=True, color=GRIS, space_after=0)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
